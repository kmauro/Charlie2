from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path('Informe_estructura_base_de_datos_Charlie2.docx')


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), color)
    tc_pr.append(shade)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    tr_pr.append(header)


def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_font(run, size=9.4, bold=False, color='222222'):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, bold=False, size=9.4, color='222222'):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f'Heading {level}')
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_font(run, size=15 if level == 1 else 11, bold=True, color='000000')
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True)
        add_text(paragraph, text[len(bold_prefix):])
    else:
        add_text(paragraph, text)
    return paragraph


def add_bullets(doc, entries):
    for entry in entries:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        add_text(p, entry)


def add_catalog(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(5.2)
    table.columns[2].width = Cm(8.0)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, value in zip(header.cells, ['Tabla', 'Finalidad', 'Claves y relaciones principales']):
        set_cell_shading(cell, '1F4E78')
        set_cell_margin(cell)
        set_cell_border(cell, 'D9D9D9')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, value, bold=True, size=8.5, color='FFFFFF')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for pos, (cell, value) in enumerate(zip(cells, row)):
            if index % 2:
                set_cell_shading(cell, 'F3F7FA')
            set_cell_margin(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, bold=(pos == 0), size=8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_field_matrix(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(7.1)
    table.columns[2].width = Cm(6.1)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, value in zip(header.cells, ['Entidad', 'Campos de negocio', 'Control e integridad']):
        set_cell_shading(cell, '385723')
        set_cell_margin(cell)
        set_cell_border(cell)
        add_text(cell.paragraphs[0], value, bold=True, size=8.5, color='FFFFFF')
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for pos, (cell, value) in enumerate(zip(cells, row)):
            if index % 2:
                set_cell_shading(cell, 'F4F8F1')
            set_cell_margin(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, bold=(pos == 0), size=8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    styles['Normal'].font.size = Pt(9.4)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(5)
    add_text(title, 'Informe de estructura de base de datos', bold=True, size=23, color='000000')
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    add_text(subtitle, 'Proyecto Charlie2 | Laravel 13 | MySQL | 28 de septiembre de 2026', size=10, color='4F4F4F')

    add_heading(doc, 'Resumen ejecutivo')
    add_body(doc, 'El sistema se apoya en una base de datos relacional MySQL orientada a la operación de una clínica: organizaciones y sedes, usuarios y profesionales, pacientes, agenda, auditoría y facturación. El diseño combina claves foráneas, restricciones de unicidad e índices compuestos para proteger la consistencia y agilizar consultas operativas.')
    add_body(doc, 'El esquema documentado proviene de 26 migraciones del repositorio. La instancia MySQL configurada no estaba disponible durante la elaboración, por lo que este informe describe el modelo que el código crea y evoluciona; no certifica las tablas efectivamente aplicadas en un servidor específico.')

    add_heading(doc, 'Alcance y convenciones')
    add_bullets(doc, [
        'Todas las tablas de dominio usan id como clave primaria incremental, salvo tablas técnicas y pivotes que usan claves propias o compuestas.',
        'timestamps representa created_at y updated_at. Las tablas de auditoría y algunas técnicas usan marcas de tiempo específicas.',
        'Restrict evita eliminar un registro referenciado; cascade elimina dependientes; nullOnDelete conserva el registro hijo y vacía la referencia.',
        'Los campos JSON almacenan colecciones o configuraciones flexibles; los campos de estado y tipo se almacenan como texto controlado por la aplicación.',
    ])

    add_heading(doc, 'Mapa funcional de relaciones')
    add_body(doc, 'La entidad legal agrupa sedes y su configuración fiscal. Cada sede concentra consultorios, feriados, caja y agenda. Los usuarios reciben roles y se vinculan con entidades legales y sedes; un usuario puede representar a un profesional o, en forma opcional, a un paciente.')
    add_body(doc, 'Las citas relacionan paciente, profesional y sede, y opcionalmente un consultorio o una serie. La agenda se completa con horarios base, pausas, excepciones, ausencias, bloqueos y lista de espera. La facturación conecta entidad legal, sede, paciente, cita, facturas, pagos, caja y contabilidad.')

    add_heading(doc, 'Catálogo de tablas')

    add_heading(doc, 'Identidad y seguridad', level=2)
    add_catalog(doc, [
        ('roles', 'Catálogo de perfiles de acceso.', 'code es único; users.role_id es opcional y restringe el borrado del rol.'),
        ('users', 'Cuentas de acceso y metadatos de gestión.', 'email e invitation_token_hash únicos; role_id → roles; status, permissions JSON e invitación.'),
        ('professionals', 'Perfil clínico asociado a una cuenta.', 'user_id único → users; documento y matrícula únicos; el borrado del usuario está restringido.'),
        ('legal_entity_user', 'Pivote de usuarios habilitados por entidad legal.', 'PK compuesta legal_entity_id + user_id; ambas FK con cascade.'),
        ('location_user', 'Pivote de usuarios habilitados por sede.', 'PK compuesta location_id + user_id; ambas FK con cascade.'),
        ('password_reset_tokens', 'Tokens para restablecer contraseñas.', 'email es clave primaria.'),
        ('sessions', 'Sesiones persistidas por Laravel.', 'id es clave primaria; user_id es opcional e indexado.'),
    ])

    add_heading(doc, 'Organización y pacientes', level=2)
    add_catalog(doc, [
        ('legal_entities', 'Personas jurídicas de la clínica.', 'CUIT único; concentra datos fiscales y estados de archivado.'),
        ('locations', 'Sedes operativas.', 'legal_entity_id → legal_entities con restrict; code único; servicios JSON y horarios.'),
        ('consulting_rooms', 'Consultorios por sede.', 'location_id → locations con cascade; único por location_id + code.'),
        ('patients', 'Historia administrativa del paciente.', 'document_number único; user_id opcional y único → users; control optimista con lock_version.'),
        ('patient_contacts', 'Contactos, responsables y emergencias.', 'patient_id → patients con cascade; índices para contacto de emergencia y responsable legal.'),
        ('audit_logs', 'Trazabilidad de operaciones.', 'actor_id opcional → users con nullOnDelete; referencia polimórfica auditable_type + auditable_id.'),
    ])

    add_heading(doc, 'Agenda y atención', level=2)
    add_catalog(doc, [
        ('professional_schedules', 'Disponibilidad semanal del profesional.', 'professional_id → professionals con cascade; location_id y consulting_room_id son opcionales.'),
        ('professional_schedule_breaks', 'Interrupciones dentro de un horario base.', 'professional_schedule_id → professional_schedules con cascade; índice por rango horario.'),
        ('professional_schedule_exceptions', 'Excepciones diarias de agenda.', 'professional_id y location_id con cascade; único por profesional, sede y fecha.'),
        ('professional_leaves', 'Licencias o ausencias del profesional.', 'professional_id obligatorio y location_id opcional; ambas FK con cascade.'),
        ('location_holidays', 'Feriados o cierres por sede.', 'location_id → locations con cascade; único por sede y fecha.'),
        ('appointments', 'Turnos clínicos.', 'patient_id, professional_id y location_id con restrict; consultorio, serie y cancelador opcionales.'),
        ('appointment_series', 'Reglas de turnos recurrentes.', 'Paciente, profesional y sede obligatorios; consultorio opcional; índices de recurrencia.'),
        ('appointment_waitlist_entries', 'Pacientes en lista de espera.', 'patient_id obligatorio; profesional, sede y confirmador opcionales.'),
        ('agenda_blocks', 'Bloqueos de agenda por alcance.', 'Puede referir sede, profesional o consultorio; created_by → users con restrict.'),
    ])

    add_heading(doc, 'Facturación y contabilidad', level=2)
    add_catalog(doc, [
        ('fiscal_configurations', 'Configuración fiscal y credenciales ARCA.', 'legal_entity_id único → legal_entities: relación uno a uno.'),
        ('accounts', 'Plan de cuentas por entidad legal.', 'Código único dentro de legal_entity_id.'),
        ('payment_methods', 'Medios de pago por entidad legal.', 'Código único dentro de legal_entity_id.'),
        ('cash_registers', 'Caja activa por sede.', 'location_id único → locations; opened_by opcional → users.'),
        ('invoices', 'Comprobantes emitidos.', 'Vincula entidad, sede, paciente, cita, creador y factura original opcionales según campo.'),
        ('invoice_items', 'Ítems de una factura.', 'invoice_id → invoices con cascade.'),
        ('installments', 'Cuotas de una factura.', 'invoice_id → invoices con cascade; número único dentro de la factura.'),
        ('payments', 'Cobros registrados.', 'Factura y caja opcionales; método de pago y receptor obligatorios.'),
        ('refunds', 'Solicitudes, autorizaciones y registros de devolución.', 'Factura y solicitante obligatorios; autorizador y registrador opcionales.'),
        ('cash_closures', 'Cierres de caja.', 'Caja, sede, entidad legal y usuario de cierre con restrict.'),
        ('journal_entries', 'Asientos contables.', 'Entidad y creador obligatorios; source_type + source_id como referencia polimórfica.'),
        ('journal_entry_lines', 'Débitos y créditos de cada asiento.', 'journal_entry_id con cascade; account_id con restrict.'),
    ])

    add_heading(doc, 'Infraestructura de aplicación', level=2)
    add_catalog(doc, [
        ('cache y cache_locks', 'Caché y bloqueos de Laravel.', 'key es clave primaria; expiration está indexado.'),
        ('jobs, job_batches y failed_jobs', 'Colas, lotes y fallos de procesos asíncronos.', 'jobs indexa queue; failed_jobs tiene uuid único e índice por conexión, cola y fecha.'),
        ('migrations', 'Historial de migraciones aplicadas por Laravel.', 'Tabla gestionada por el framework al ejecutar las migraciones.'),
    ])

    add_heading(doc, 'Campos de negocio relevantes')
    add_field_matrix(doc, [
        ('Usuarios y profesionales', 'users: name, email, status, permissions, último acceso e invitación. professionals: documento, especialidad, matrícula, duración de turno y asignación de consultorio.', 'Email, documento y usuario profesional son únicos. La matrícula compuesta también es única.'),
        ('Pacientes', 'Documento, nombres, nacimiento, sexo registrado, contacto, cobertura, afiliado, estado y fecha de archivo.', 'Documento único. user_id es único y opcional. lock_version y previous_status apoyan la concurrencia.'),
        ('Sedes y consultorios', 'Nombre, código, dirección, barrio, horario, zona horaria, servicios JSON y estado. Consultorios: código, sector, piso, equipamiento JSON y restricciones.', 'CUIT y código de sede son únicos. El código de consultorio es único dentro de la sede.'),
        ('Turnos', 'Inicio, fin, duración, tipo, estado, motivo, sobrecupo, cancelación, recursos y preparación del procedimiento.', 'Índices por profesional, paciente, sede, consultorio y rango horario. lock_version evita actualizaciones perdidas.'),
        ('Facturas y cobros', 'Tipo, estado, número, moneda, importes neto, impuestos, total y pagado. Ítems, cuotas, pagos, devoluciones y cierres incluyen importes monetarios decimal 14,2.', 'Número único por entidad legal. Los ítems y cuotas se eliminan al eliminar la factura.'),
        ('Auditoría y contabilidad', 'Audit logs guarda antes y después, evento, resultado, razón, correlación, IP y agente. Asientos guardan fecha, descripción, origen y líneas de débito/crédito.', 'Auditoría y origen contable usan pares tipo + id para referencias polimórficas.'),
    ])

    add_heading(doc, 'Reglas de integridad y rendimiento')
    add_body(doc, 'El modelo privilegia restrictOnDelete en entidades maestras y transacciones importantes: roles, usuarios profesionales, entidades legales, sedes, pacientes, profesionales, turnos, cajas, facturas, pagos y cuentas. Esta decisión impide perder historial cuando existen registros dependientes.')
    add_body(doc, 'Se utiliza cascadeOnDelete para información subordinada que no tiene sentido sin su padre: contactos de paciente, franjas y pausas de agenda, feriados, algunas relaciones de sede, ítems de factura, cuotas y líneas de asiento. En referencias históricas no obligatorias se prefiere nullOnDelete, por ejemplo actor de auditoría, paciente de factura, usuario que abrió caja y responsables de la lista de espera.')
    add_body(doc, 'Los índices compuestos se concentran en patrones de consulta operativa: agenda por profesional, paciente, sede y consultorio; búsqueda de pacientes activos por nombre; disponibilidad de profesionales; feriados y excepciones; espera por prioridad; facturación por sede y estado; cierres por entidad y sede. Las unicidades protegen identificadores de negocio y evitan duplicar horarios, códigos, accesos de pivote o números de comprobante.')

    add_heading(doc, 'Observaciones operativas')
    add_bullets(doc, [
        'La configuración fiscal guarda certificado y clave privada en campos de texto. Debe protegerse con cifrado de aplicación, controles de acceso y rotación de secretos.',
        'Los campos JSON ofrecen flexibilidad para permisos, servicios, equipamiento y recursos de procedimientos; las consultas frecuentes sobre sus contenidos deben evaluarse antes de incorporar índices especializados.',
        'La disponibilidad real surge de combinar horario base, pausas, excepciones, licencias, feriados, bloqueos y turnos existentes; no depende de una única tabla.',
        'Para comprobar que una instalación concreta coincide con este informe se debe ejecutar el estado de migraciones y consultar el esquema de MySQL cuando el servidor esté disponible.',
    ])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, 'Informe de estructura de base de datos | Charlie2', size=8, color='666666')

    doc.core_properties.title = 'Informe de estructura de base de datos Charlie2'
    doc.core_properties.subject = 'Documentación del esquema de base de datos'
    doc.core_properties.author = 'Charlie2'
    doc.save(OUTPUT)


if __name__ == '__main__':
    main()
